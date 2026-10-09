from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

from config.settings import settings


class ExcelReportGenerator:
    """Genera el reporte diario de alarmas activas en formato Excel."""

    NARANJA = "E8722C"
    AMARILLO = "F4B41A"
    GRIS_OSC = "3A3A3A"
    GRIS_FILA = "F2F2F2"
    BLANCO = "FFFFFF"
    ROJO = "E05252"

    HEADERS = ["Hora de alarma", "Nombre de la planta", "SN", "Equipo",
               "Tipo de alarma", "Nombre de la alarma", "Nivel", "Estado",
               "Duración", "Sugerencia"]

    def __init__(self, alarmas, logo_path=None):
        """alarmas: filas obtenidas de AlarmaRepository.obtener_activas()."""
        self.alarmas = alarmas
        self.logo_path = logo_path or (settings.base_dir / "assets" / "logo_sunnyapp.png")
        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active
        self.ws.title = "Reporte diario"

    def generar(self) -> str:
        self._encabezado()
        self._tabla()
        self._resumen()
        self._ajustar_anchos()

        nombre_archivo = f"reporte_diario_{datetime.now().strftime('%Y%m%d')}.xlsx"
        ruta = settings.excel_dir / nombre_archivo
        self.wb.save(ruta)
        return str(ruta)

    def _encabezado(self):
        if self.logo_path.exists():
            img = XLImage(str(self.logo_path))
            img.width, img.height = 70, 73
            self.ws.add_image(img, "A1")

        self.ws.merge_cells("B1:E1")
        self.ws["B1"] = "SUNNY APP"
        self.ws["B1"].font = Font(size=18, bold=True, color=self.NARANJA)

        self.ws.merge_cells("B2:E2")
        self.ws["B2"] = "Reporte diario de alarmas — Plataforma Goodwe"
        self.ws["B2"].font = Font(size=12, bold=True, color=self.GRIS_OSC)

        self.ws.merge_cells("F1:J1")
        self.ws["F1"] = "Fecha del reporte:"
        self.ws["F1"].font = Font(bold=True, size=10)
        self.ws["F1"].alignment = Alignment(horizontal="right")

        self.ws.merge_cells("F2:J2")
        self.ws["F2"] = datetime.now().strftime("%d/%m/%Y")
        self.ws["F2"].alignment = Alignment(horizontal="right")

        for r in (1, 2, 3):
            self.ws.row_dimensions[r].height = 22

    def _tabla(self):
        header_row = 5
        for i, h in enumerate(self.HEADERS, start=1):
            c = self.ws.cell(row=header_row, column=i, value=h)
            c.font = Font(bold=True, color=self.BLANCO, size=10)
            c.fill = PatternFill("solid", fgColor=self.GRIS_OSC)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        self.ws.row_dimensions[header_row].height = 28

        thin = Side(style="thin", color="D9D9D9")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for r_i, alarma in enumerate(self.alarmas, start=header_row + 1):
            valores = [
                alarma["hora_alarma"], alarma["nombre_planta"], alarma["sn"], alarma["equipo"],
                alarma["tipo_alarma"], alarma["nombre_alarma"], alarma["nivel"], alarma["estado"],
                alarma["duracion"] or "-", alarma["sugerencia"] or "-",
            ]
            for c_i, val in enumerate(valores, start=1):
                c = self.ws.cell(row=r_i, column=c_i, value=val)
                c.font = Font(size=10)
                c.border = border
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if r_i % 2 == 0:
                    c.fill = PatternFill("solid", fgColor=self.GRIS_FILA)

            nivel_cell = self.ws.cell(row=r_i, column=7)
            nivel_cell.fill = PatternFill("solid", fgColor=self.ROJO if nivel_cell.value == "Fallo" else self.AMARILLO)
            nivel_cell.font = Font(bold=True, color=self.BLANCO)

            estado_cell = self.ws.cell(row=r_i, column=8)
            estado_cell.font = Font(bold=True, color="C0392B" if estado_cell.value == "Ocurriendo" else "2E7D5B")

        self._ultima_fila = header_row + len(self.alarmas)
        self.ws.freeze_panes = "A6"

    def _resumen(self):
        resumen_row = self._ultima_fila + 2
        self.ws.merge_cells(f"A{resumen_row}:C{resumen_row}")
        self.ws[f"A{resumen_row}"] = "Resumen del día"
        self.ws[f"A{resumen_row}"].font = Font(bold=True, size=11, color=self.NARANJA)

        total = len(self.alarmas)
        fallos = sum(1 for a in self.alarmas if a["nivel"] == "Fallo")

        self.ws[f"A{resumen_row + 1}"] = "Total alarmas activas"
        self.ws[f"B{resumen_row + 1}"] = total
        self.ws[f"A{resumen_row + 2}"] = "Fallos críticos"
        self.ws[f"B{resumen_row + 2}"] = fallos

        for rr in (resumen_row + 1, resumen_row + 2):
            self.ws[f"A{rr}"].font = Font(size=10)
            self.ws[f"B{rr}"].font = Font(size=10, bold=True)

    def _ajustar_anchos(self):
        widths = [18, 26, 18, 10, 18, 26, 10, 12, 12, 30]
        for i, w in enumerate(widths, start=1):
            self.ws.column_dimensions[get_column_letter(i)].width = w